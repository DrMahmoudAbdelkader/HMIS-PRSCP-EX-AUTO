"""
queue_mr_extractor  -  v13  (DIRECT appointment-dashboard extraction)
=====================================================================
Replaces the v12 WebReport/Excel pathway. It no longer imports
queue_extractor.py or queue_parser.py - you can delete both from the repo.

WHAT IT DOES
------------
Talks to the HMIS appointment dashboard the same way the browser does
(JSF/PrimeFaces partial-ajax calls recorded in queue_direct_extraction.har):

  1. Launch the appointment module      POST  /HMIS/            (login)
  2. Open the appointment clerk page    GET   .../appointmentClerk.xhtml
  3. Pick the day                       click the calendar day link
  4. For EVERY speciality in the dropdown (all clinics, not only daycare):
        change specDropDown -> reload -> click each doctor's getSlotBtn
  5. Parse every slot: time, patient name, MR, status (Booking / Arrived /
     No Show / ...), doctor, clinic.
  6. Keep ONLY slots whose status is an "arrived" status. No-shows and
     not-yet-arrived bookings drop out here (no arrival-date column needed).

PUBLIC API  (unchanged - Extract_DMS_..._modified.py needs no edits)
-------------------------------------------------------------------
    get_queue_records(run_date_ddmmyyyy, output_dir=None) -> list[dict]
    get_queue_mr_clinic_map(run_date_ddmmyyyy, output_dir=None) -> {MR: Clinic}
        RAISES on a real failure (never sys.exit).

Audit workbook (every slot incl. no-shows/bookings + per-clinic counts):
    <output_dir>/<dd-mm-yyyy>_queue_direct.xlsx

ENVIRONMENT
-----------
  HMIS_QUEUE_USERNAME / HMIS_QUEUE_PASSWORD   dashboard login (REQUIRED;
                                              NOT the CMIS account)
  QUEUE_OUTPUT_DIR         audit folder (default D:\\Queue_DMS_Data)
  QUEUE_ARRIVED_STATUSES   default "arrived,seen"
  QUEUE_EXCLUDE_SPECS      speciality codes/names to skip, default none
                           (e.g. "17,DP,24" = archive + pharmacies)
  QUEUE_DELAY              seconds between calls (default 0.3)

CLI
---
  python queue_mr_extractor.py                # today
  python queue_mr_extractor.py 28-09-2026     # a given day
"""

import base64
import html as html_lib
import os
import re
import sys
import time
import warnings
from collections import Counter
from datetime import datetime

import requests
from bs4 import BeautifulSoup

warnings.filterwarnings("ignore", message=".*XMLParsedAsHTMLWarning.*")
try:
    from bs4 import XMLParsedAsHTMLWarning
    warnings.filterwarnings("ignore", category=XMLParsedAsHTMLWarning)
except Exception:
    pass

# ======================================================================
# CONFIG
# ======================================================================
HOST        = "41.33.24.254:8080"
BASE        = f"http://{HOST}"
HMIS        = f"{BASE}/HMIS"
CLERK_PAGE  = f"{HMIS}/faces/appointment/appointmentClerk.xhtml"
CLERK_LINK  = f"{HMIS}/faces/appointment/appointmentClerkPage.xhtml"
MENU_LOGIN  = f"{BASE}/hmis-menu/faces/management/login.xhtml"

HOSPITAL_B64 = base64.b64encode(b"01").decode()    # encb
BRANCH_B64   = base64.b64encode(b"1").decode()     # encbid
LANG_B64     = base64.b64encode(b"L").decode()     # encl
MODULE_B64   = base64.b64encode(b"APN").decode()   # encm (appointment module)

# Copied verbatim from the recorded module-launch request.
CATEGORY_MODULE_LIST = "[{\"categoryCode\":\"1\",\"categoryEnName\":\"Hospital\",\"categoryArName\":\"\u0645\u0633\u062a\u0634\u0641\u064a\",\"categoryName\":\"Hospital\",\"categoryColor\":\"bg-orange\",\"moduleCode\":null,\"moduleName\":null,\"moduleUrl\":null,\"moduleIcon\":null,\"moduleList\":[{\"moduleName\":null,\"moduleCode\":\"BAS\",\"encryptedModuleName\":\"QkFT\",\"loginUrl\":\"http://41.33.24.254:8080/HMIS/\",\"openWeb\":false,\"iconName\":\"new/inpatient.png\",\"colorBlock\":null,\"descAr\":\"\u0627\u0644\u0645\u0631\u0636\u0649 \u0627\u0644\u062f\u0627\u062e\u0644\u0649\",\"descEn\":\"Inpatient Management\",\"ipAddress\":\" \",\"port\":\" \",\"addConfig\":\"N\"},{\"moduleName\":null,\"moduleCode\":\"PCY\",\"encryptedModuleName\":\"UENZ\",\"loginUrl\":\"http://41.33.24.254:8080/HmisPharmacy/\",\"openWeb\":false,\"iconName\":\"new/hospitals/pharmacy-icon.png\",\"colorBlock\":null,\"descAr\":\"\u0627\u0644\u0635\u064a\u062f\u0644\u064a\u0629\",\"descEn\":\"Pharmacy\",\"ipAddress\":\" \",\"port\":\" \",\"addConfig\":\"N\"},{\"moduleName\":null,\"moduleCode\":\"PMI\",\"encryptedModuleName\":\"UE1J\",\"loginUrl\":\"http://41.33.24.254:8080/HMIS/\",\"openWeb\":false,\"iconName\":\"registeration-icon.png\",\"colorBlock\":null,\"descAr\":\"\u0627\u0644\u062a\u0633\u062c\u064a\u0644\",\"descEn\":\"Registration\",\"ipAddress\":\" \",\"port\":\" \",\"addConfig\":\"N\"},{\"moduleName\":null,\"moduleCode\":\"APN\",\"encryptedModuleName\":\"QVBO\",\"loginUrl\":\"http://41.33.24.254:8080/HMIS/\",\"openWeb\":false,\"iconName\":\"new/Out-Patient.png\",\"colorBlock\":null,\"descAr\":\"\u0645\u0648\u0627\u0639\u064a\u062f \u0627\u0644\u0639\u064a\u0627\u062f\u0627\u062a \u0644\u0644\u0645\u0631\u0636\u0649 \u0627\u0644\u062e\u0627\u0631\u062c\u064a\",\"descEn\":\"Outpatient\",\"ipAddress\":\" \",\"port\":\" \",\"addConfig\":\"N\"},{\"moduleName\":null,\"moduleCode\":\"ROP\",\"encryptedModuleName\":\"Uk9Q\",\"loginUrl\":\"http://41.33.24.254:8080/HMIS/\",\"openWeb\":false,\"iconName\":\"dashboard.png\",\"colorBlock\":null,\"descAr\":\"\u062a\u0642\u0631\u064a\u0631 \u0627\u0644\u0639\u0645\u0644\u064a\u0627\u062a\",\"descEn\":\"Operation Report\",\"ipAddress\":\" \",\"port\":\" \",\"addConfig\":\"N\"},{\"moduleName\":null,\"moduleCode\":\"EMR\",\"encryptedModuleName\":\"RU1S\",\"loginUrl\":\"http://41.33.24.253/CMIS/MRM_5.2/login/sec\",\"openWeb\":false,\"iconName\":\"new/medical-records.png\",\"colorBlock\":null,\"descAr\":\"\u0627\u0644\u0633\u062c\u0644\u0627\u062a \u0627\u0644\u0637\u0628\u064a\u0629 \u0627\u0644\u0625\u0644\u0643\u062a\u0631\u0648\u0646\u064a\u0629\",\"descEn\":\"EMR\",\"ipAddress\":\" \",\"port\":\" \",\"addConfig\":\"N\"},{\"moduleName\":null,\"moduleCode\":\"ORE\",\"encryptedModuleName\":\"T1JF\",\"loginUrl\":\"http://41.33.24.254:8080/HMIS/\",\"openWeb\":false,\"iconName\":\"orders-icon.png\",\"colorBlock\":null,\"descAr\":\"\u0625\u062f\u062e\u0627\u0644 \u0627\u0644\u0637\u0644\u0628\",\"descEn\":\"Order Entry\",\"ipAddress\":\" \",\"port\":\" \",\"addConfig\":\"N\"}]},{\"categoryCode\":\"4\",\"categoryEnName\":\"Material Management\",\"categoryArName\":\"\u0625\u062f\u0627\u0631\u0647 \u0627\u0644\u0645\u0648\u0627\u062f\",\"categoryName\":\"Material Management\",\"categoryColor\":\"bg-red\",\"moduleCode\":null,\"moduleName\":null,\"moduleUrl\":null,\"moduleIcon\":null,\"moduleList\":[{\"moduleName\":null,\"moduleCode\":\"INV\",\"encryptedModuleName\":\"SU5W\",\"loginUrl\":\"http://41.33.24.254:59360/DMS-INV-war/faces/index_new.xhtml\",\"openWeb\":false,\"iconName\":\"inventory.png\",\"colorBlock\":null,\"descAr\":\"\u0646\u0638\u0627\u0645 \u0623\u062f\u0627\u0631\u0629 \u0627\u0644\u0645\u062e\u0632\u0648\u0646\",\"descEn\":\"Inventory Management System\",\"ipAddress\":\" \",\"port\":\" \",\"addConfig\":\"N\"}]}]"

OUTPUT_DIR = os.environ.get("QUEUE_OUTPUT_DIR", r"D:\Queue_DMS_Data")
TIMEOUT    = 60
RETRIES    = 3
DELAY      = float(os.environ.get("QUEUE_DELAY", "0.3"))

ARRIVED_STATUSES = {s.strip().lower() for s in
                    os.environ.get("QUEUE_ARRIVED_STATUSES", "arrived,seen").split(",") if s.strip()}
EXCLUDE_SPECS = {s.strip().lower() for s in
                 os.environ.get("QUEUE_EXCLUDE_SPECS", "").split(",") if s.strip()}
KNOWN_STATUSES = {"booking", "arrived", "seen", "no show", "noshow", "cancelled", "canceled"}

_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
       "(KHTML, like Gecko) Chrome/150.0.0.0 Safari/537.36")

_MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}

# Exact partial-render lists the browser sends (from the HAR).
RENDER_DAY  = ("appointmentClerkForm menuForm:viewDate menuForm:specDropDown "
               "menuForm:resourceDropDown menuForm:doctorDropDown menuForm:clinicDropDown "
               "menuForm:reloadBtn appointmentClerkForm:selectedSearchDate menuForm:patientDetails")
RENDER_SPEC = ("menuForm:specDropDown menuForm:resourceDropDown menuForm:doctorDropDown "
               "menuForm:clinicDropDown menuForm:reloadBtn")


def _warn(msg):
    print(f"   !! {msg}")
    if os.environ.get("GITHUB_ACTIONS") == "true":
        print("::warning::" + str(msg).replace("\r", " ").replace("\n", " "))


def _s(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


# ======================================================================
# CLIENT
# ======================================================================
class QueueClient:
    def __init__(self, session=None):
        self.s = session or requests.Session()
        self.s.headers.update({"User-Agent": _UA,
                               "Accept-Language": "en-US,en;q=0.9,ar;q=0.8"})
        self.viewstate = ""
        self.form_html = ""          # latest <appointmentClerkForm> fragment
        self.spec = ""               # currently selected speciality code
        self.display_mode = "2"      # 2 = "all"
        self.cal_html = ""           # latest calendar fragment (day links)
        self.date_label = ""
        self.first_page = ""

    # -- low level ----------------------------------------------------
    def _req(self, method, url, **kw):
        last = None
        for attempt in range(1, RETRIES + 1):
            try:
                r = self.s.request(method, url, timeout=TIMEOUT, **kw)
                if r.status_code >= 500:
                    raise RuntimeError(f"HTTP {r.status_code}")
                return r
            except Exception as e:
                last = e
                print(f"   !! {method} {url} try {attempt}/{RETRIES} failed: {e}")
                time.sleep(1.5 * attempt)
        raise RuntimeError(f"{method} {url} failed after {RETRIES} tries: {last}")

    @staticmethod
    def _updates(xml_text):
        """{update id: html} from a JSF partial-response."""
        return {m.group(1): m.group(2) for m in re.finditer(
            r'<update id="([^"]+)"><!\[CDATA\[(.*?)\]\]></update>', xml_text, re.S)}

    def _partial(self, params):
        params = dict(params)
        params["javax.faces.partial.ajax"] = "true"
        params["javax.faces.ViewState"] = self.viewstate
        r = self._req("POST", CLERK_PAGE, data=params, headers={
            "Faces-Request": "partial/ajax", "X-Requested-With": "XMLHttpRequest",
            "Accept": "application/xml, text/xml, */*; q=0.01",
            "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
            "Origin": BASE, "Referer": CLERK_PAGE})
        txt = r.text
        if "<redirect" in txt or "login.xhtml" in txt[:2000]:
            raise RuntimeError("HMIS session expired / redirected to login")
        if "<error>" in txt:
            msg = re.search(r"<error-message>(.*?)</error-message>", txt, re.S)
            raise RuntimeError("JSF error: " + _s(msg.group(1) if msg else txt[:300]))
        ups = self._updates(txt)
        vs = ups.get("j_id1:javax.faces.ViewState:0")
        if vs:
            self.viewstate = vs
        if "appointmentClerkForm" in ups:
            self.form_html = ups["appointmentClerkForm"]
        if "menuForm:viewDate" in ups:
            m = re.search(r"<h2>([^<]+)</h2>", ups["menuForm:viewDate"])
            if m:
                self.date_label = _s(m.group(1))
        if "menuForm:calenderController" in ups:
            self.cal_html = ups["menuForm:calenderController"]
        time.sleep(DELAY)
        return ups

    def _menu_fields(self, spec=None):
        return {"menuForm": "menuForm",
                "menuForm:displayMode": self.display_mode,
                "menuForm:specDropDown": self.spec if spec is None else spec,
                "menuForm:resourceDropDown": "",
                "menuForm:doctorDropDown": "-1",
                "menuForm:clinicDropDown": ""}

    # -- login --------------------------------------------------------
    def login(self, username=None, password=None):
        username = username or os.environ.get("HMIS_QUEUE_USERNAME", "")
        password = password or os.environ.get("HMIS_QUEUE_PASSWORD", "")
        if not username or not password:
            raise RuntimeError("HMIS_QUEUE_USERNAME / HMIS_QUEUE_PASSWORD are not set")
        try:                                    # seed JSESSIONID (best effort)
            self.s.get(MENU_LOGIN, timeout=TIMEOUT)
        except Exception:
            pass
        self.s.cookies.set("dmsmodule", "SRV", domain=HOST.split(":")[0], path="/")
        r = self._req("POST", f"{HMIS}/", data={
            "encu": base64.b64encode(username.encode()).decode(),
            "encp": base64.b64encode(password.encode()).decode(),
            "encb": HOSPITAL_B64, "encbid": BRANCH_B64, "encl": LANG_B64,
            "CategoryModuleList": CATEGORY_MODULE_LIST, "encm": MODULE_B64,
        }, headers={"Origin": BASE, "Referer": MENU_LOGIN,
                    "Content-Type": "application/x-www-form-urlencoded"})
        if r.status_code != 200:
            raise RuntimeError(f"HMIS launch failed: HTTP {r.status_code}")
        self._req("GET", CLERK_LINK, headers={"Referer": f"{HMIS}/"})
        page = self._req("GET", CLERK_PAGE, headers={"Referer": CLERK_LINK}).text
        self._load_first_page(page)
        print("   -> HMIS login OK, appointment page loaded")

    def _load_first_page(self, page):
        if "menuForm:specDropDown" not in page:
            raise RuntimeError("Appointment page did not load (login rejected or "
                               "module not authorised). First 200 chars: " + _s(page[:200]))
        vs = re.search(r'name="javax\.faces\.ViewState" id="j_id1:javax\.faces\.ViewState:0" value="([^"]+)"', page)
        if not vs:
            raise RuntimeError("Could not find the JSF ViewState on the appointment page")
        self.viewstate = html_lib.unescape(vs.group(1))
        self.first_page = page
        self.cal_html = page
        m = re.search(r'id="menuForm:viewDate">\s*<h2>([^<]+)</h2>', page)
        self.date_label = _s(m.group(1)) if m else ""

    # -- calendar -----------------------------------------------------
    def _calendar_month(self):
        m = re.search(r'<div class="month">\s*([A-Za-z]+)\.?\s+(\d{4})\s*</div>', self.cal_html)
        if not m:
            return None
        return int(m.group(2)), _MONTHS.get(m.group(1)[:3].lower())

    def _calendar_cells(self):
        """[(anchor_id, day_number)] for every numeric calendar anchor, in page order.
        Tolerant on purpose: JSF auto-generated ids (j_idt104 / j_idt106 / j_idt108 ...)
        change when the page structure changes, so we do not rely on them."""
        cells = []
        for m in re.finditer(r'<a\b[^>]*?\bid="([^"]+)"[^>]*>(.*?)</a>', self.cal_html, re.S):
            aid = m.group(1)
            if not aid.startswith("menuForm:"):
                continue
            txt = html_lib.unescape(re.sub(r"<[^>]+>", "", m.group(2))).strip()
            if txt.isdigit() and 1 <= int(txt) <= 31:
                cells.append((aid, int(txt)))
        return cells

    def _day_link_id(self, day):
        cells = self._calendar_cells()
        # The grid may also show trailing days of the previous month (28,29,30 ...)
        # before the 1st and leading days of the next month after the last day.
        # Keep only the increasing run that starts at the first "1".
        start = next((i for i, (_, n) in enumerate(cells) if n == 1), 0)
        run, prev = [], 0
        for aid, n in cells[start:]:
            if n <= prev:
                break
            run.append((aid, n))
            prev = n
        for aid, n in run:
            if n == day:
                return aid
        return None

    def _calendar_diag(self):
        cells = self._calendar_cells()
        print(f"   !! calendar diag: {len(self.cal_html)} chars, "
              f"{len(cells)} numeric day anchors, month={self._calendar_month()}")
        print("   !! day anchors (first 45): " +
              ", ".join(f"{n}={a.replace('menuForm:', '')}" for a, n in cells[:45]))
        snippet = re.sub(r"\s+", " ", self.cal_html)
        i = snippet.find('class="month"')
        print("   !! calendar html near month header: " +
              snippet[max(0, i - 100): i + 1500] if i >= 0 else
              "   !! no month header; html start: " + snippet[:1500])

    def _step_month(self, direction):
        cls = "glyphicon-chevron-left" if direction < 0 else "glyphicon-chevron-right"
        m = re.search(r'<a id="([^"]+)"[^>]*class="glyphicon %s[^"]*"' % cls, self.cal_html)
        if not m:
            raise RuntimeError("calendar chevron not found")
        src = m.group(1)
        self._partial({"javax.faces.source": src, "javax.faces.partial.execute": src,
                       "javax.faces.partial.render": "menuForm:calenderController",
                       "javax.faces.behavior.event": "click", "javax.faces.partial.event": "click",
                       **self._menu_fields()})

    def select_date(self, date_dash):
        d = datetime.strptime(date_dash, "%d-%m-%Y")
        want = (d.year, d.month)
        for _ in range(24):
            cur = self._calendar_month()
            if cur is None or cur == want:
                break
            self._step_month(-1 if want < cur else 1)
        link = self._day_link_id(d.day)
        if not link:
            # Current month: no chevron click happened, so cal_html is still the
            # INITIAL full page, whose calendar markup differs from the partial
            # update that works for past months. Force a partial re-render by
            # stepping one month away and back.
            try:
                self._step_month(-1)
                self._step_month(+1)
            except Exception as e:
                print(f"   !! calendar re-render failed: {e}")
            link = self._day_link_id(d.day)
        if not link:
            self._calendar_diag()
            try:                      # keep the raw calendar HTML for diagnosis
                os.makedirs(OUTPUT_DIR, exist_ok=True)
                with open(os.path.join(OUTPUT_DIR, f"calendar_debug_{date_dash}.html"),
                          "w", encoding="utf-8") as fh:
                    fh.write(self.cal_html)
            except Exception:
                pass
            raise RuntimeError(f"No calendar day link for {date_dash} "
                               f"(calendar month shown: {self._calendar_month()})")
        self.display_mode = "1"
        self._partial({"javax.faces.source": link, "javax.faces.partial.execute": link,
                       "javax.faces.partial.render": RENDER_DAY,
                       "javax.faces.behavior.event": "click", "javax.faces.partial.event": "click",
                       **self._menu_fields(spec="")})
        self.display_mode = "2"
        self.spec = ""
        got = self.selected_date()
        if got != d.strftime("%Y-%m-%d"):
            raise RuntimeError(f"Dashboard is showing {got!r}, not {d:%Y-%m-%d} - "
                               f"refusing to extract the wrong day (header: {self.date_label!r})")

    def selected_date(self):
        m = re.search(r'name="appointmentClerkForm:selectedSearchDate"[^>]*value="([^"]*)"', self.form_html) \
            or re.search(r'value="([^"]*)"[^>]*name="appointmentClerkForm:selectedSearchDate"', self.form_html) \
            or re.search(r'id="appointmentClerkForm:selectedSearchDate"[^>]*value="([^"]*)"', self.form_html)
        if not m:
            return None
        try:
            return datetime.strptime(html_lib.unescape(m.group(1)).strip(), "%b %d, %Y").strftime("%Y-%m-%d")
        except ValueError:
            return None

    # -- specialities -------------------------------------------------
    def list_specialities(self):
        m = re.search(r'<select[^>]*id="menuForm:specDropDown".*?</select>', self.first_page, re.S)
        out = []
        for code, label in re.findall(r'<option value="([^"]*)"[^>]*>(.*?)</option>', m.group(0), re.S):
            if not code:
                continue
            label = _s(html_lib.unescape(re.sub(r"<[^>]+>", "", label)))
            out.append((code, re.sub(r"^\S+\s*:\s*", "", label)))
        return out

    def load_speciality(self, code):
        """Select the speciality, reload, return [(doc_index, doctor_name)]."""
        self.spec = code
        self._partial({"javax.faces.source": "menuForm:specDropDown",
                       "javax.faces.partial.execute": "menuForm:specDropDown",
                       "javax.faces.partial.render": RENDER_SPEC,
                       "javax.faces.behavior.event": "valueChange", "javax.faces.partial.event": "change",
                       **self._menu_fields()})
        self._partial({"javax.faces.source": "menuForm:reloadBtn",
                       "javax.faces.partial.execute": "menuForm",
                       "javax.faces.partial.render": "appointmentClerkForm",
                       "javax.faces.behavior.event": "click", "javax.faces.partial.event": "click",
                       **self._menu_fields()})
        return self._doctors()

    def _doctors(self):
        docs = []
        soup = BeautifulSoup(self.form_html, "lxml")
        for blk in soup.find_all(id=re.compile(r"^appointmentClerkForm:specLoop:0:docLoop:\d+:slotAction$")):
            idx = int(re.search(r"docLoop:(\d+):", blk["id"]).group(1))
            t = blk.select_one(".booking-doctor .dropdown-toggle")
            docs.append((idx, _s(t.get_text(" ")) if t else f"doctor {idx}"))
        return docs

    # -- slots --------------------------------------------------------
    def _form_payload(self):
        """Every named <input> of appointmentClerkForm, as the browser serialises it."""
        soup = BeautifulSoup(self.form_html, "lxml")
        pay = {"appointmentClerkForm": "appointmentClerkForm"}
        for inp in soup.find_all("input"):
            n = inp.get("name")
            if n and n.startswith("appointmentClerkForm:") and inp.get("type") not in (
                    "button", "submit", "checkbox", "radio", "file"):
                pay[n] = html_lib.unescape(inp.get("value", ""))
        return pay

    def load_doctor_slots(self, doc_idx):
        src = f"appointmentClerkForm:specLoop:0:docLoop:{doc_idx}:getSlotBtn"
        self._partial({"javax.faces.source": src, "javax.faces.partial.execute": src,
                       "javax.faces.partial.render": "appointmentClerkForm",
                       "javax.faces.behavior.event": "click", "javax.faces.partial.event": "click",
                       **self._form_payload()})
        return parse_slots(self.form_html, spec_loop=0, doc_loop=doc_idx)


def parse_slots(form_html, spec_loop=None, doc_loop=None):
    """Parse patient slots out of an appointmentClerkForm fragment."""
    soup = BeautifulSoup("<div>" + form_html + "</div>", "lxml")
    rows = []
    for div in soup.select('div[id$=":slot"]'):
        m = re.match(r"appointmentClerkForm:specLoop:(\d+):docLoop:(\d+):slotSection:(\d+):slot$", div.get("id", ""))
        if not m:
            continue
        sp, dc, sl = (int(x) for x in m.groups())
        if (spec_loop is not None and sp != spec_loop) or (doc_loop is not None and dc != doc_loop):
            continue
        pat = div.select_one(".slot-patient")
        labels = [_s(l.get_text()) for l in pat.find_all("label")] if pat else []
        if len(labels) < 2 or not labels[1]:
            continue                                  # empty (free) slot
        st = div.select_one("div.status span.dropdown-toggle")
        proc = pat.select_one(".hasprocedure")
        t = div.select_one(".slot-time")
        rows.append({"spec_loop": sp, "doc_loop": dc, "slot_no": sl,
                     "Slot Time": _s(t.get_text()) if t else "",
                     "Patient Name": labels[0], "MR": labels[1],
                     "Status": _s(st.get_text()) if st else "",
                     "Procedure": _s(proc.get("title")) if proc else ""})
    return rows


# ======================================================================
# ORCHESTRATION
# ======================================================================
def _is_arrived(status):
    return _s(status).lower() in ARRIVED_STATUSES


def extract_day(date_dash, client=None):
    """All slots for one day, all specialities."""
    c = client or QueueClient()
    if client is None:
        c.login()
    print(f"   -> Selecting {date_dash}")
    c.select_date(date_dash)
    want_iso = datetime.strptime(date_dash, "%d-%m-%Y").strftime("%Y-%m-%d")
    specs = [(code, name) for code, name in c.list_specialities()
             if code.lower() not in EXCLUDE_SPECS and name.lower() not in EXCLUDE_SPECS]
    print(f"   -> {len(specs)} speciality(ies): " + ", ".join(n for _, n in specs))
    out, seen = [], set()
    for code, name in specs:
        docs = c.load_speciality(code)
        if c.selected_date() not in (None, want_iso):
            raise RuntimeError("dashboard date changed while extracting")
        n_spec = 0
        for doc_idx, doc_name in docs:
            for r in c.load_doctor_slots(doc_idx):
                key = (code, doc_name, r["Slot Time"], r["MR"])
                if key in seen:
                    continue
                seen.add(key)
                r.update({"Spec Code": code, "Clinic": name, "Doctor": doc_name,
                          "Appointment Date": date_dash})
                out.append(r)
                n_spec += 1
        cnt = Counter(_s(r["Status"]) for r in out if r["Spec Code"] == code)
        print(f"      {name:<24} doctors={len(docs):<2} slots={n_spec:<4} {dict(cnt)}")
    return out


def _write_audit(rows, date_dash, out_dir):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"{date_dash}_queue_direct.xlsx")
    cols = ["Appointment Date", "Clinic", "Doctor", "Slot Time", "MR", "Patient Name",
            "Status", "Procedure", "Arrived?"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Arrived"
    ws2 = wb.create_sheet("All slots")
    ws3 = wb.create_sheet("By clinic")
    for w in (ws, ws2):
        w.append(cols)
        for c in w[1]:
            c.font = Font(bold=True)
        w.freeze_panes = "A2"
    for r in rows:
        vals = [r.get(c, "") for c in cols[:-1]] + ["Y" if _is_arrived(r["Status"]) else "N"]
        ws2.append(vals)
        if _is_arrived(r["Status"]):
            ws.append(vals)
    statuses = sorted({_s(r["Status"]) for r in rows})
    ws3.append(["Clinic"] + statuses + ["Total", "Arrived"])
    for c in ws3[1]:
        c.font = Font(bold=True)
    for clinic in sorted({r["Clinic"] for r in rows}):
        rr = [r for r in rows if r["Clinic"] == clinic]
        cnt = Counter(_s(r["Status"]) for r in rr)
        ws3.append([clinic] + [cnt.get(s, 0) for s in statuses]
                   + [len(rr), sum(_is_arrived(r["Status"]) for r in rr)])
    wb.save(path)
    return path


def get_queue_records(run_date_ddmmyyyy, output_dir=None):
    out_dir = output_dir if output_dir is not None else OUTPUT_DIR
    rows = extract_day(run_date_ddmmyyyy)
    unknown = {_s(r["Status"]) for r in rows if _s(r["Status"]).lower() not in KNOWN_STATUSES}
    if unknown:
        _warn(f"Unrecognised slot status(es) seen: {sorted(unknown)} - "
              f"add to QUEUE_ARRIVED_STATUSES if they mean 'patient arrived'")
    try:
        p = _write_audit(rows, run_date_ddmmyyyy, out_dir)
        print(f"   -> Saved audit workbook -> {p}")
    except Exception as e:
        _warn(f"could not write audit workbook: {e}")
    arrived = [r for r in rows if _is_arrived(r["Status"])]
    print(f"   -> {len(rows)} booked slot(s), {len(arrived)} arrived, "
          f"{len({r['MR'] for r in arrived})} unique MR")
    return [{
        "MR": r["MR"], "Clinic": r["Clinic"], "National ID": "", "Appointment Number": "",
        "Appointment Date": r["Appointment Date"], "Serial No.": "", "Status": r["Status"],
        "User": "", "Old Medical No.": "", "Source": "HMIS appointmentClerk (direct)",
        "Doctor": r["Doctor"], "Slot Time": r["Slot Time"], "Patient Name": r["Patient Name"],
    } for r in arrived]


def get_queue_mr_clinic_map(run_date_ddmmyyyy, output_dir=None):
    out, conflicts = {}, []
    for r in get_queue_records(run_date_ddmmyyyy, output_dir):
        mr = r["MR"]
        if mr in out:
            if out[mr] != r["Clinic"]:
                conflicts.append((mr, out[mr], r["Clinic"]))
            continue
        out[mr] = r["Clinic"]
    for mr, kept, dropped in conflicts[:10]:
        print(f"   !! MR {mr} arrived at several clinics: kept {kept!r}, also saw {dropped!r}")
    return out


if __name__ == "__main__":
    day = sys.argv[1] if len(sys.argv) > 1 else datetime.now().strftime("%d-%m-%Y")
    recs = get_queue_records(day)
    print(f"{len(recs)} arrived rows, {len({r['MR'] for r in recs})} unique MR")
